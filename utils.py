# -*- coding: utf-8 -*-
"""Business logic utilities for HR PMI System."""

import re
import csv
import io
import difflib
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

WORKFORCE_HEADERS = ["사번(익명ID)", "생년", "나이", "직급", "직급순서", "입사년도", "근속년수", "소속", "직무구분", "관리자여부"]
EMPLOYEE_HEADERS = ["사번", "성명", "생년월일", "주소", "소속", "직급", "고용형태", "입사일", "월급여", "계약시작일", "계약종료일", "승계여부", "특이사항"]
REG_STOPWORDS = {
    "한다", "하여", "대한", "관한", "따라", "경우", "있다", "없다", "위하여", "또는", "그리고", "다만",
    "이하", "각", "등", "및", "의", "을", "를", "이", "가", "은", "는", "에", "에서", "으로", "로",
    "제조", "조의", "회사", "직원", "기관", "규정", "시행", "적용", "정한다", "아니한다"
}


def read_text_flexible(content: bytes) -> str:
    for enc in ["utf-8", "utf-8-sig", "cp949", "euc-kr", "utf-16"]:
        try:
            return content.decode(enc)
        except Exception:
            pass
    return content.decode(errors="ignore")


def extract_pdf(content: bytes) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            pages.append(f"[PDF Page {i}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages).strip()


def extract_docx(content: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(content))
    out = []
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text.strip())
    for table in d.tables:
        for row in table.rows:
            vals = [c.text.strip() for c in row.cells]
            if any(vals):
                out.append(" | ".join(vals))
    return "\n".join(out).strip()


def extract_hwpx(content: bytes) -> str:
    out = []
    with zipfile.ZipFile(io.BytesIO(content), "r") as z:
        names = [n for n in z.namelist() if n.lower().endswith(".xml")]
        targets = [n for n in names if "section" in n.lower()] or names
        for name in targets:
            try:
                root = ET.fromstring(z.read(name))
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        out.append(elem.text.strip())
            except Exception:
                continue
    return "\n".join(out).strip()


def read_any_file(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        return read_text_flexible(content)
    if ext == ".pdf":
        return extract_pdf(content)
    if ext == ".docx":
        return extract_docx(content)
    if ext == ".hwpx":
        return extract_hwpx(content)
    try:
        return read_text_flexible(content)
    except Exception:
        raise ValueError(f"지원하지 않는 파일 형식: {ext}")


def read_xlsx_rows(content: bytes):
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(io.BytesIO(content), "r") as z:
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.findall(f"{{{ns_main}}}si"):
                shared.append("".join(t.text or "" for t in si.iter(f"{{{ns_main}}}t")))
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        sheet = wb.find(f".//{{{ns_main}}}sheet")
        if sheet is None:
            return []
        rel_id = sheet.attrib.get(f"{{{ns_rel}}}id", "")
        target = ""
        for rel in rels.findall(f"{{{ns_pkg}}}Relationship"):
            if rel.attrib.get("Id") == rel_id:
                target = rel.attrib.get("Target", "")
                break
        if not target:
            return []
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target.lstrip("./")
        root = ET.fromstring(z.read(target))
        rows = []
        for row_elem in root.findall(f".//{{{ns_main}}}row"):
            row = []
            for cell in row_elem.findall(f"{{{ns_main}}}c"):
                ref = cell.attrib.get("r", "A1")
                letters = re.match(r"[A-Z]+", ref)
                col = 0
                for ch in (letters.group(0) if letters else "A"):
                    col = col * 26 + ord(ch) - 64
                col -= 1
                while len(row) <= col:
                    row.append("")
                cell_type = cell.attrib.get("t", "")
                value_elem = cell.find(f"{{{ns_main}}}v")
                value = value_elem.text if value_elem is not None and value_elem.text is not None else ""
                if cell_type == "s" and value:
                    try:
                        value = shared[int(value)]
                    except Exception:
                        pass
                elif cell_type == "inlineStr":
                    value = "".join(t.text or "" for t in cell.iter(f"{{{ns_main}}}t"))
                row[col] = value
            rows.append(row)
        return rows


def read_csv_bytes(content: bytes):
    for enc in ["utf-8-sig", "cp949", "euc-kr", "utf-8"]:
        try:
            text = content.decode(enc)
            return list(csv.reader(io.StringIO(text)))
        except Exception:
            pass
    raise RuntimeError("CSV 파일을 읽지 못했습니다.")


def clean_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_articles(text: str):
    text = clean_text(text)
    pat = re.compile(r"(?=(?:^|\n)\s*제\s*\d+\s*조(?:의\s*\d+)?\s*(?:\([^)\n]*\))?)")
    parts = [p.strip() for p in pat.split(text) if p.strip()]
    if len(parts) <= 1:
        return [{"key": f"문단-{i+1}", "text": l.strip()} for i, l in enumerate(text.splitlines()) if l.strip()]
    rows = []
    for i, p in enumerate(parts):
        m = re.match(r"\s*(제\s*\d+\s*조(?:의\s*\d+)?)", p)
        key = re.sub(r"\s+", "", m.group(1)) if m else f"기타-{i+1}"
        rows.append({"key": key, "text": p})
    return rows


def article_map(arts):
    d = {}
    for a in arts:
        key = a["key"]
        if key in d:
            n = 2
            while f"{key}-{n}" in d:
                n += 1
            key = f"{key}-{n}"
        d[key] = a
    return d


def natural_key(s):
    nums = re.findall(r"\d+", s)
    return [int(nums[0]) if nums else 999999, int(nums[1]) if len(nums) > 1 else 0, s]


def similarity(a, b):
    return difflib.SequenceMatcher(None, a, b).ratio()


def risk_tag(text):
    s = text or ""
    if re.search(r"보수|임금|수당|성과급|복지|휴가|연차|정년|징계|해고|면직|승진|전보|근무지|노조|단체협약|퇴직", s):
        return "노무·불이익변경 검토"
    if re.search(r"채용|평가|시험|면접|전형", s):
        return "공정채용·평가절차 검토"
    return "일반검토"


def compare_regulations(a_text, b_text):
    ma = article_map(split_articles(a_text))
    mb = article_map(split_articles(b_text))
    rows = []
    for key in sorted(set(ma) | set(mb), key=natural_key):
        a = ma.get(key)
        b = mb.get(key)
        if a and not b:
            rows.append(["A에만 있음", key, "A사규에만 존재합니다.", risk_tag(a["text"]), a["text"], ""])
        elif b and not a:
            rows.append(["B에만 있음", key, "B사규에만 존재합니다.", risk_tag(b["text"]), "", b["text"]])
        else:
            r = similarity(a["text"], b["text"])
            if r >= 0.985:
                rows.append(["동일", key, "동일", "-", a["text"], b["text"]])
            else:
                rows.append(["차이있음", key, "문구·내용 차이가 있습니다.", risk_tag(a["text"] + b["text"]), a["text"], b["text"]])
    return rows


def regulation_words(text):
    tokens = re.findall(r"[가-힣A-Za-z0-9]{2,}", clean_text(text).lower())
    result = {}
    for token in tokens:
        token = re.sub(r"^(제|각)", "", token)
        if len(token) < 2 or token in REG_STOPWORDS or token.isdigit():
            continue
        result[token] = result.get(token, 0) + 1
    return result


def find_duplicate_regulations(a_text, b_text, threshold=0.58):
    wa, wb = regulation_words(a_text), regulation_words(b_text)
    common_words = sorted(
        [[word, wa[word], wb[word], wa[word] + wb[word]] for word in set(wa) & set(wb)],
        key=lambda row: (-row[3], row[0])
    )
    arts_a, arts_b = split_articles(a_text), split_articles(b_text)
    similar = []
    for aa in arts_a:
        norm_a = re.sub(r"\s+", "", aa["text"])
        if len(norm_a) < 12:
            continue
        candidates = []
        for bb in arts_b:
            norm_b = re.sub(r"\s+", "", bb["text"])
            if len(norm_b) < 12:
                continue
            score = similarity(norm_a, norm_b)
            if score >= threshold:
                candidates.append((score, bb))
        for score, bb in sorted(candidates, key=lambda x: x[0], reverse=True)[:3]:
            similar.append([
                aa["key"], bb["key"], round(score * 100, 1),
                aa["text"][:700], bb["text"][:700]
            ])
    similar.sort(key=lambda row: (-row[2], natural_key(row[0])))
    return common_words, similar


def to_number(value):
    try:
        return float(str(value).replace(",", "").strip() or 0)
    except Exception:
        return 0.0


def workforce_stats(workforce, base_year):
    if not workforce:
        return {}
    normalized = []
    for row in workforce:
        row = (list(row) + [""] * len(WORKFORCE_HEADERS))[:len(WORKFORCE_HEADERS)]
        birth = to_number(row[1])
        row[2] = str(int(base_year - birth)) if birth else row[2]
        normalized.append(row)
    n = len(normalized)
    avg_age = sum(to_number(r[2]) for r in normalized) / n
    avg_tenure = sum(to_number(r[6]) for r in normalized) / n
    grades = sorted(set(to_number(r[4]) for r in normalized))
    min_grade, max_grade = grades[0], grades[-1]
    span = max(1.0, max_grade - min_grade)
    bands = {"상위": 0, "중간": 0, "하위": 0}
    for row in normalized:
        pos = (to_number(row[4]) - min_grade) / span
        bands["상위" if pos <= .3 else "하위" if pos >= .7 else "중간"] += 1
    mid_share = bands["중간"] / n
    managers = sum(1 for r in normalized if str(r[9]).strip().upper() in ("Y", "YES", "1", "관리자"))
    if mid_share >= .4 and bands["중간"] > bands["상위"] and bands["중간"] > bands["하위"]:
        shape = "항아리형"
    elif bands["하위"] > bands["중간"]:
        shape = "피라미드형"
    elif bands["상위"] > bands["중간"]:
        shape = "역피라미드형"
    else:
        shape = "균형형"
    return {
        "n": n, "avg_age": avg_age, "avg_tenure": avg_tenure,
        "mid_share": mid_share, "manager_share": managers / n,
        "bands": bands, "shape": shape, "normalized": normalized
    }


def workforce_group(workforce, index, classifier=None):
    counts = {}
    for row in workforce:
        key = classifier(row) if classifier else str(row[index])
        counts[key] = counts.get(key, 0) + 1
    return counts


def normalize_workforce_rows(rows):
    header_idx = next((i for i, row in enumerate(rows) if any(str(v).strip() == "직급" for v in row)), -1)
    if header_idx < 0:
        raise RuntimeError("헤더 행에서 '직급' 열을 찾지 못했습니다.")
    headers = [str(v).strip() for v in rows[header_idx]]
    required = ["생년", "직급", "직급순서", "입사년도", "근속년수"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise RuntimeError("필수 열 누락: " + ", ".join(missing))
    result = []
    for source in rows[header_idx + 1:]:
        if not any(str(v).strip() for v in source):
            continue
        item = {h: (source[i] if i < len(source) else "") for i, h in enumerate(headers)}
        result.append([item.get(h, "") for h in WORKFORCE_HEADERS])
    return result


def normalize_employee_rows(rows):
    normalized = []
    for row in rows or []:
        row = list(row)
        if len(row) == 8:
            row = ["", row[0], "", "", row[1], row[2], row[3], row[4], row[5], "", "", row[6], row[7]]
        normalized.append((row + [""] * len(EMPLOYEE_HEADERS))[:len(EMPLOYEE_HEADERS)])
    return normalized


def task_rate(tasks):
    if not tasks:
        return 0
    weights = {"대기": 0, "진행": 50, "완료": 100}
    return round(sum(weights.get(t[4] if len(t) > 4 else "", 0) for t in tasks) / len(tasks))


def generate_contract_docx(template_bytes, mapping):
    import docx
    doc = docx.Document(io.BytesIO(template_bytes))
    replacements = {}
    for key, value in mapping.items():
        value = str(value or "")
        for placeholder in (f"{{{{{key}}}}}", f"${{{key}}}", f"[{key}]", f"«{key}»"):
            replacements[placeholder] = value
    hits = 0
    for paragraph in doc.paragraphs:
        original = "".join(run.text for run in paragraph.runs)
        updated = original
        for ph, val in replacements.items():
            updated = updated.replace(ph, val)
        if updated != original:
            hits += 1
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    original = "".join(run.text for run in paragraph.runs)
                    updated = original
                    for ph, val in replacements.items():
                        updated = updated.replace(ph, val)
                    if updated != original:
                        hits += 1
                        if paragraph.runs:
                            paragraph.runs[0].text = updated
                            for run in paragraph.runs[1:]:
                                run.text = ""
    if hits == 0:
        doc.add_paragraph("")
        doc.add_heading("개인별 계약정보", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key in EMPLOYEE_HEADERS:
            r = table.add_row().cells
            r[0].text = key
            r[1].text = str(mapping.get(key, ""))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), hits


def create_sample_contract_template():
    import docx
    doc = docx.Document()
    doc.add_heading("근로계약서(고용승계)", level=1)
    doc.add_paragraph("{{통합기관명}}(이하 '회사')와 {{성명}}(이하 '근로자')는 다음과 같이 근로계약을 체결한다.")
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key in ["사번", "성명", "생년월일", "주소", "소속", "직급", "고용형태", "입사일", "월급여", "계약시작일", "계약종료일"]:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = "{{" + key + "}}"
    doc.add_paragraph("작성일: {{작성일}}")
    doc.add_paragraph("회사: {{통합기관명}}    근로자: {{성명}} (서명)")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_report_docx(text):
    import docx
    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def rows_to_csv_bytes(rows):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8-sig")
